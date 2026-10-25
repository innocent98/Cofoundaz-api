"""Documents BE-C, generator and legal templates (gap-list rows 13 and 18)."""

import io

import pytest
from docx import Document as DocxDocument
from reportlab import rl_config
from sqlalchemy.orm import Session

from app.core.config import settings
from app.db.models.business import BusinessCanvas, BusinessRecord
from app.db.models.document import Document, DocumentVersion
from app.db.models.enums import DOCUMENT_ROLES, CanvasType, MembershipRole, RecordKind
from app.db.models.job import Job
from app.db.models.startup import Startup
from app.platform.llm_budget import debit
from app.services.documents.template_defs import DISCLAIMER
from tests.documents_helpers import headers, team

DENIED = sorted(set(MembershipRole) - DOCUMENT_ROLES, key=lambda r: r.value)
LEGAL = ["nda", "contractor_agreement", "founder_agreement", "safe"]


@pytest.fixture(autouse=True)
def _env(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "STORAGE_BACKEND", "local")
    monkeypatch.setattr(settings, "LOCAL_STORAGE_DIR", str(tmp_path))
    monkeypatch.setattr(settings, "LLM_PROVIDER", "stub")
    monkeypatch.setattr(rl_config, "pageCompression", 0)


@pytest.fixture
def ws(db: Session):
    startup, users = team(db)
    startup.industry = "Fintech"
    startup.description = "Payments for market traders."
    startup.country = "NG"
    db.flush()
    return startup, users, {r: headers(u, startup) for r, u in users.items()}


def _seed_builder(db: Session, startup: Startup) -> None:
    db.add(
        BusinessCanvas(
            startup_id=startup.id,
            type=CanvasType.lean,
            blocks={
                "problem": ["Cash handling is slow", "Fraud"],
                "solution": ["Mobile wallet"],
                "customer_segments": ["Market traders"],
            },
        )
    )
    db.add(
        BusinessRecord(
            startup_id=startup.id,
            kind=RecordKind.revenue_stream,
            data={"name": "Transaction fees", "pricing_basis": "1% per payment"},
            position=0,
        )
    )
    db.flush()


def _form(client, h, key):
    r = client.get(f"/api/v1/document-templates/{key}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()["data"]


def _inst(client, h, key, **body):
    return client.post(f"/api/v1/documents/templates/{key}/instantiate", json=body, headers=h)


def _section(doc, heading):
    return next(s for s in doc["sections"] if s["heading"] == heading)


# ---- catalog ---------------------------------------------------------------------------------


def test_catalog_uses_the_four_prd_categories_and_adds_the_legal_drafts(client, ws):
    _s, _u, heads = ws
    r = client.get("/api/v1/document-templates", headers=heads[MembershipRole.founder])
    templates = r.json()["data"]["templates"]
    assert {t["category"] for t in templates} == {"Legal", "Financial", "Pitch", "Operational"}
    by_key = {t["key"]: t for t in templates}
    assert {"business_plan", "pitch_deck", "financial_model", "meeting_notes", "one_pager"} <= set(
        by_key
    )
    for key in LEGAL:
        assert by_key[key]["category"] == "Legal"
        assert by_key[key]["disclaimer"] == DISCLAIMER
        assert by_key[key]["kind"] == key
    assert by_key["business_plan"]["disclaimer"] is None
    assert by_key["financial_model"]["category"] == "Financial"
    assert by_key["meeting_notes"]["category"] == "Operational"
    assert isinstance(by_key["nda"]["sections"][0], str)  # the old shape (headings) is kept


def test_disclaimer_says_not_legal_advice():
    assert "not legal advice" in DISCLAIMER.lower()


# ---- form state --------------------------------------------------------------------------------


def test_form_state_marks_autofilled_empty_and_ai_draftable(client, ws, db):
    startup, _u, heads = ws
    _seed_builder(db, startup)
    form = _form(client, heads[MembershipRole.founder], "business_plan")["form"]
    secs = {s["heading"]: s for s in form["sections"]}
    assert secs["Problem"]["state"] == "filled" and secs["Problem"]["source"] == "business_builder"
    assert secs["Problem"]["value"] == "- Cash handling is slow\n- Fraud"
    assert secs["Executive Summary"]["source"] == "profile"
    assert secs["Executive Summary"]["value"] == "Payments for market traders."
    assert secs["Business Model"]["value"].startswith("- Transaction fees")
    assert "1% per payment" in secs["Business Model"]["value"]
    assert secs["Go-to-Market"]["state"] == "empty" and secs["Go-to-Market"]["ai_draftable"]
    assert secs["Go-to-Market"]["source"] is None
    assert secs["The Ask"]["state"] == "empty" and secs["The Ask"]["ai_draftable"] is False
    assert all(s["key"] for s in form["sections"])
    assert form["ai_available"] is True


def test_form_state_without_builder_data_is_mostly_empty(client, ws):
    _s, _u, heads = ws
    form = _form(client, heads[MembershipRole.team_member], "business_plan")["form"]
    secs = {s["heading"]: s for s in form["sections"]}
    assert secs["Problem"]["state"] == "empty"
    assert secs["Executive Summary"]["state"] == "filled"  # the profile description is enough


def test_legal_form_prefills_party_and_date_from_the_account_and_leaves_the_rest(client, ws):
    _s, _u, heads = ws
    form = _form(client, heads[MembershipRole.founder], "nda")["form"]
    fields = {f["key"]: f for f in form["fields"]}
    assert fields["disclosing_party"]["value"] == "Acme"
    assert fields["disclosing_party"]["source"] == "profile"
    assert fields["effective_date"]["state"] == "filled"
    assert fields["effective_date"]["source"] == "system"
    assert fields["receiving_party"]["state"] == "empty" and fields["receiving_party"]["required"]
    assert fields["receiving_party"]["ai_draftable"] is False
    assert form["unfilled_required"] == [
        k for k, f in fields.items() if f["required"] and f["state"] == "empty"
    ]
    assert form["disclaimer"] == DISCLAIMER
    assert all(s["ai_draftable"] is False for s in form["sections"])
    assert form["ai_available"] is False


@pytest.mark.parametrize("role", DENIED, ids=lambda r: r.value)
def test_other_roles_cannot_read_the_form(client, ws, role):
    _s, _u, heads = ws
    assert client.get("/api/v1/document-templates/nda", headers=heads[role]).status_code == 403


def test_unknown_template_is_404(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    assert client.get("/api/v1/document-templates/nope", headers=h).status_code == 404
    assert _inst(client, h, "nope").status_code == 404


# ---- instantiate -------------------------------------------------------------------------------


def test_instantiate_creates_the_document_with_autofill_and_returns_the_form(client, ws, db):
    startup, _u, heads = ws
    _seed_builder(db, startup)
    h = heads[MembershipRole.founder]
    r = _inst(client, h, "business_plan", title="NovaPay plan", folder="Fundraising")
    assert r.status_code == 201, r.text
    data = r.json()["data"]
    doc = data["document"]
    assert doc["title"] == "NovaPay plan" and doc["folder"] == "Fundraising"
    assert doc["template_key"] == "business_plan" and doc["kind"] == "business_plan"
    assert doc["ai_generated"] is False and doc["version"] == 1
    assert _section(doc, "Problem")["body"] == "- Cash handling is slow\n- Fraud"
    assert _section(doc, "Go-to-Market")["body"] == ""
    assert data["form"]["sections"][0]["state"] in ("filled", "empty")
    assert data["ai"] == {"status": "not_requested", "job_id": None, "sections": []}
    assert db.query(Job).filter(Job.type == "ai.documents.template_draft").count() == 0
    got = client.get(f"/api/v1/documents/{doc['id']}", headers=h).json()["data"]
    assert got["sections"] == doc["sections"]


def test_team_member_can_instantiate_and_default_title_is_the_template_name(client, ws):
    _s, _u, heads = ws
    r = _inst(client, heads[MembershipRole.team_member], "one_pager")
    assert r.status_code == 201
    assert r.json()["data"]["document"]["title"] == "One-Pager"


def test_legal_template_fills_values_and_marks_the_gaps(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    r = _inst(
        client,
        h,
        "nda",
        values={"receiving_party": "Bright Ltd", "governing_law": "Lagos State, Nigeria"},
    )
    assert r.status_code == 201, r.text
    data = r.json()["data"]
    doc = data["document"]
    text = " ".join(s["body"] for s in doc["sections"])
    assert "Bright Ltd" in text and "Lagos State, Nigeria" in text and "Acme" in text
    assert "{{" not in text
    assert doc["kind"] == "nda" and doc["disclaimer"] == DISCLAIMER
    fields = {f["key"]: f for f in data["form"]["fields"]}
    assert fields["receiving_party"]["source"] == "you"
    assert fields["receiving_party"]["state"] == "filled"
    assert "receiving_party" not in data["form"]["unfilled_required"]
    # a field left empty shows up as a visible bracket, never a silent blank
    r2 = _inst(client, h, "nda")
    text2 = " ".join(s["body"] for s in r2.json()["data"]["document"]["sections"])
    assert "[Receiving party]" in text2


@pytest.mark.parametrize("key", LEGAL)
def test_each_legal_template_is_a_structured_draft_with_the_disclaimer(client, ws, key):
    _s, _u, heads = ws
    data = _inst(client, heads[MembershipRole.founder], key).json()["data"]
    doc = data["document"]
    assert len(doc["sections"]) >= 6
    assert all(s["body"].strip() for s in doc["sections"])
    assert doc["disclaimer"] == DISCLAIMER and doc["kind"] == key
    assert any(f["required"] for f in data["form"]["fields"])


def test_summary_rows_also_carry_the_disclaimer_flag_for_legal_kinds(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    _inst(client, h, "safe")
    rows = client.get("/api/v1/documents", headers=h).json()["data"]["documents"]
    assert rows[0]["disclaimer"] == DISCLAIMER and rows[0]["type"] == "document"
    _inst(client, h, "one_pager")
    rows = client.get("/api/v1/documents", headers=h).json()["data"]["documents"]
    assert {r["template_key"]: r["disclaimer"] for r in rows}["one_pager"] is None


def test_post_documents_with_a_legal_template_gets_the_draft_text(client, ws):
    _s, _u, heads = ws
    r = client.post(
        "/api/v1/documents", json={"template_key": "safe"}, headers=heads[MembershipRole.founder]
    )
    assert r.status_code == 201
    assert len(r.json()["data"]["sections"]) >= 6


def test_unknown_field_and_oversized_value_are_422(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    bad = _inst(client, h, "nda", values={"nonsense": "x"})
    assert bad.status_code == 422
    assert bad.json()["error"]["field_errors"][0]["field"] == "values.nonsense"
    assert _inst(client, h, "nda", values={"purpose": "x" * 5001}).status_code == 422
    assert _inst(client, h, "nda", values={"receiving_party": 5}).status_code == 422


def test_control_characters_in_values_are_stripped(client, ws):
    _s, _u, heads = ws
    r = _inst(
        client,
        heads[MembershipRole.founder],
        "nda",
        values={"receiving_party": "Bright\x00 Ltd\x07"},
    )
    text = " ".join(s["body"] for s in r.json()["data"]["document"]["sections"])
    assert "Bright Ltd" in text and "\x00" not in text and "\x07" not in text


def test_placeholder_syntax_in_a_value_is_not_expanded(client, ws):
    _s, _u, heads = ws
    r = _inst(
        client,
        heads[MembershipRole.founder],
        "nda",
        values={"receiving_party": "{{disclosing_party}}", "governing_law": "Nigeria"},
    )
    text = " ".join(s["body"] for s in r.json()["data"]["document"]["sections"])
    assert "{{disclosing_party}}" in text  # inserted literally, not re-expanded


@pytest.mark.parametrize("role", DENIED, ids=lambda r: r.value)
def test_other_roles_cannot_instantiate(client, ws, db, role):
    _s, _u, heads = ws
    before = db.query(Document).count()
    assert _inst(client, heads[role], "nda").status_code == 403
    assert db.query(Document).count() == before


def test_legal_exports_print_the_disclaimer(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "nda").json()["data"]["document"]

    def fetch(fmt):
        f = client.post(
            f"/api/v1/documents/{doc['id']}/export", json={"format": fmt}, headers=h
        ).json()["data"]
        link = client.get(f"/api/v1/documents/files/{f['id']}/download", headers=h).json()["data"]
        return client.get("/api/v1" + link["url"].split("/api/v1", 1)[1]).content

    pdf = fetch("pdf").decode("latin-1")
    assert "not legal advice" in pdf.replace("\\", "")
    parsed = DocxDocument(io.BytesIO(fetch("docx")))
    assert any(DISCLAIMER == p.text for p in parsed.paragraphs)
    footer = " ".join(p.text for s in parsed.sections for p in s.footer.paragraphs)
    assert "not legal advice" in footer.lower()


# ---- AI drafts ---------------------------------------------------------------------------------


def _run(db, job_id):
    from app.worker.handlers.documents_ai import handle_template_draft

    job = db.get(Job, job_id)
    job.attempts = 1
    handle_template_draft(db, job)
    db.flush()
    return job


def test_ai_draft_enqueues_a_job_and_the_worker_fills_only_empty_ai_sections(client, ws, db):
    startup, _u, heads = ws
    _seed_builder(db, startup)
    h = heads[MembershipRole.founder]
    r = _inst(client, h, "business_plan", ai_draft=True)
    assert r.status_code == 201, r.text
    data = r.json()["data"]
    ai = data["ai"]
    assert ai["status"] == "queued" and ai["job_id"]
    assert "go_to_market" in ai["sections"] and "problem" not in ai["sections"]
    assert "the_ask" not in ai["sections"] and "team" not in ai["sections"]
    job = db.get(Job, ai["job_id"])
    assert job.type == "ai.documents.template_draft" and job.startup_id == startup.id
    doc_id = data["document"]["id"]

    # the founder edits one of the sections the AI was going to write, before the worker runs
    doc = db.get(Document, doc_id)
    secs = [dict(s) for s in doc.sections]
    next(s for s in secs if s["heading"] == "Financials")["body"] = "My own numbers"
    doc.sections = secs
    db.flush()

    _run(db, job.id)
    db.refresh(job)
    assert job.result["status"] == "ready"
    assert "financials" in job.result["skipped"]
    assert "go_to_market" in job.result["drafted"]
    db.refresh(doc)
    by_h = {s["heading"]: s["body"] for s in doc.sections}
    assert "[stub-llm]" in by_h["Go-to-Market"]
    assert by_h["Financials"] == "My own numbers"
    assert by_h["Problem"] == "- Cash handling is slow\n- Fraud"  # autofill untouched
    assert by_h["The Ask"] == ""
    assert doc.ai_generated is True and doc.version == 2
    assert "stub-llm" in doc.search_text
    snap = db.query(DocumentVersion).filter_by(document_id=doc.id, version=2).one()
    assert snap.author_id is None and "AI" in snap.note


def test_ai_draft_is_refused_when_the_template_has_nothing_to_draft(client, ws):
    _s, _u, heads = ws
    r = _inst(client, heads[MembershipRole.founder], "nda", ai_draft=True)
    assert r.status_code == 422
    assert "AI" in r.json()["error"]["message"]


def test_ai_sections_must_be_draftable(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    r = _inst(client, h, "business_plan", ai_draft=True, ai_sections=["the_ask"])
    assert r.status_code == 422
    assert r.json()["error"]["field_errors"][0]["field"] == "ai_sections"
    ok = _inst(client, h, "business_plan", ai_draft=True, ai_sections=["go_to_market"])
    assert ok.json()["data"]["ai"]["sections"] == ["go_to_market"]


def test_over_budget_at_request_time_still_creates_the_document_and_says_why(client, ws, db):
    startup, _u, heads = ws
    debit(db, startup.id, settings.LLM_DAILY_TOKEN_BUDGET + 1)
    r = _inst(client, heads[MembershipRole.founder], "business_plan", ai_draft=True)
    assert r.status_code == 201
    data = r.json()["data"]
    assert data["ai"] == {"status": "over_budget", "job_id": None, "sections": []}
    assert data["document"]["id"]
    assert db.query(Job).filter(Job.type == "ai.documents.template_draft").count() == 0


def test_over_budget_when_the_worker_runs_writes_nothing_and_says_so(client, ws, db):
    startup, _u, heads = ws
    h = heads[MembershipRole.founder]
    data = _inst(client, h, "business_plan", ai_draft=True).json()["data"]
    debit(db, startup.id, settings.LLM_DAILY_TOKEN_BUDGET + 1)
    job = _run(db, data["ai"]["job_id"])
    assert job.result["status"] == "over_budget" and job.result["drafted"] == []
    doc = db.get(Document, data["document"]["id"])
    assert doc.version == 1 and doc.ai_generated is False


def test_llm_failure_retries_then_fails_visibly_on_the_last_attempt(client, ws, db, monkeypatch):
    startup, _u, heads = ws
    data = _inst(client, heads[MembershipRole.founder], "business_plan", ai_draft=True).json()[
        "data"
    ]
    from app.platform import llm_budget
    from app.worker.handlers import documents_ai

    class Boom:
        last_usage_tokens = 0

        def complete(self, *a, **k):
            raise RuntimeError("provider down")

    monkeypatch.setattr(llm_budget, "get_llm_client", lambda: Boom())
    job = db.get(Job, data["ai"]["job_id"])
    job.attempts = 1
    with pytest.raises(RuntimeError):
        documents_ai.handle_template_draft(db, job)  # not the last attempt: the runner retries
    job.attempts = settings.WORKER_MAX_ATTEMPTS
    documents_ai.handle_template_draft(db, job)
    assert job.result["status"] == "failed" and job.result["drafted"] == []
    assert job.result["reason"] == "generation_failed"
    assert db.get(Document, data["document"]["id"]).version == 1


def test_partial_result_is_saved_when_the_last_attempt_fails_midway(client, ws, db, monkeypatch):
    _s, _u, heads = ws
    data = _inst(client, heads[MembershipRole.founder], "business_plan", ai_draft=True).json()[
        "data"
    ]
    from app.platform import llm_budget
    from app.worker.handlers import documents_ai

    class Flaky:
        last_usage_tokens = 0
        calls = 0

        def complete(self, *a, **k):
            Flaky.calls += 1
            if Flaky.calls > 2:
                raise RuntimeError("rate limited")
            return "Drafted text"

    monkeypatch.setattr(llm_budget, "get_llm_client", lambda: Flaky())
    job = db.get(Job, data["ai"]["job_id"])
    job.attempts = settings.WORKER_MAX_ATTEMPTS
    documents_ai.handle_template_draft(db, job)
    assert job.result["status"] == "partial" and len(job.result["drafted"]) == 2
    assert len(job.result["failed"]) >= 1
    doc = db.get(Document, data["document"]["id"])
    assert doc.version == 2 and sum("Drafted text" in s["body"] for s in doc.sections) == 2


def test_worker_skips_a_trashed_or_missing_document(client, ws, db):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    data = _inst(client, h, "business_plan", ai_draft=True).json()["data"]
    client.delete(f"/api/v1/documents/{data['document']['id']}", headers=h)
    job = _run(db, data["ai"]["job_id"])
    assert job.result["status"] == "gone"
    job.payload = {**job.payload, "document_id": "00000000-0000-0000-0000-000000000000"}
    from app.worker.handlers.documents_ai import handle_template_draft

    handle_template_draft(db, job)
    assert job.result["status"] == "gone"


def test_job_is_readable_by_a_member_through_the_jobs_api(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    data = _inst(client, h, "business_plan", ai_draft=True).json()["data"]
    r = client.get(f"/api/v1/jobs/{data['ai']['job_id']}", headers=h)
    assert r.status_code == 200 and r.json()["data"]["status"] == "queued"


# ---- retry route -------------------------------------------------------------------------------


def test_redraft_route_queues_a_new_job_for_the_still_empty_sections(client, ws, db):
    _s, _u, heads = ws
    h = heads[MembershipRole.team_member]
    doc = _inst(client, heads[MembershipRole.founder], "business_plan").json()["data"]["document"]
    r = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert r.status_code == 202, r.text
    ai = r.json()["data"]
    assert ai["status"] == "queued" and "go_to_market" in ai["sections"]
    _run(db, ai["job_id"])
    after = client.get(f"/api/v1/documents/{doc['id']}", headers=h).json()["data"]
    assert "[stub-llm]" in _section(after, "Go-to-Market")["body"]
    assert "the_ask" not in ai["sections"]  # sections the AI cannot know are never queued
    again = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert again.status_code == 422  # everything it may draft is now written


def test_redraft_with_nothing_left_is_422_and_blank_documents_are_422(client, ws, db):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "nda").json()["data"]["document"]
    r = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert r.status_code == 422
    blank = client.post("/api/v1/documents", json={"title": "x"}, headers=h).json()["data"]
    assert (
        client.post(f"/api/v1/documents/{blank['id']}/ai-draft", json={}, headers=h).status_code
        == 422
    )


def test_redraft_roles_trash_and_over_budget(client, ws, db):
    startup, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "business_plan").json()["data"]["document"]
    for role in DENIED:
        assert (
            client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=heads[role])
        ).status_code == 403
    debit(db, startup.id, settings.LLM_DAILY_TOKEN_BUDGET + 1)
    over = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert over.status_code == 429 and over.json()["error"]["code"] == "AI_BUDGET_EXHAUSTED"
    client.delete(f"/api/v1/documents/{doc['id']}", headers=h)
    assert (
        client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    ).status_code == 404


def test_redraft_rejects_non_draftable_section_keys(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "business_plan").json()["data"]["document"]
    r = client.post(
        f"/api/v1/documents/{doc['id']}/ai-draft", json={"sections": ["the_ask"]}, headers=h
    )
    assert r.status_code == 422


def test_a_second_redraft_while_one_is_in_flight_is_409(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "business_plan").json()["data"]["document"]
    assert (
        client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h).status_code
        == 202
    )
    again = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert again.status_code == 409 and again.json()["error"]["code"] == "AI_DRAFT_IN_PROGRESS"


def test_redraft_ignores_sections_the_user_deleted(client, ws):
    _s, _u, heads = ws
    h = heads[MembershipRole.founder]
    doc = _inst(client, h, "one_pager").json()["data"]["document"]
    keep = [s for s in doc["sections"] if s["heading"] not in ("Overview", "Problem", "Solution")]
    client.put(
        f"/api/v1/documents/{doc['id']}",
        json={"title": doc["title"], "sections": keep, "version": doc["version"]},
        headers=h,
    )
    r = client.post(f"/api/v1/documents/{doc['id']}/ai-draft", json={}, headers=h)
    assert r.status_code == 422
