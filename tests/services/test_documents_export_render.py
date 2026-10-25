"""The renderers and template helpers behind export and the generator (no HTTP)."""

import io
import re
from types import SimpleNamespace

import pytest
from docx import Document as DocxDocument
from reportlab import rl_config

from app.db.models.enums import DocumentKind
from app.services.documents import export
from app.services.documents.generator import Context, _record_line, build_form, resolve_spec
from app.services.documents.template_defs import (
    DISCLAIMER,
    DOCUMENT_TEMPLATES,
    disclaimer_for,
    instantiate,
    render_body,
)


def _doc(sections, title="T", kind=DocumentKind.custom, template_key=None):
    return SimpleNamespace(
        title=title, version=3, sections=sections, kind=kind, template_key=template_key
    )


@pytest.fixture(autouse=True)
def _plain_pdf(monkeypatch):
    monkeypatch.setattr(rl_config, "pageCompression", 0)


def test_blocks_split_paragraphs_bullets_and_keep_single_newlines():
    got = export.blocks("Intro\nline two\n\n- a\n* b\n• c\n\nOutro\x00")
    assert got == [("p", "Intro\nline two"), ("li", "a"), ("li", "b"), ("li", "c"), ("p", "Outro")]
    assert export.blocks("") == []


def test_filename_is_safe_and_never_empty():
    assert export._safe_filename('Plan: "v2"/final?', ".pdf") == "Plan v2 final.pdf"
    assert export._safe_filename("   ", ".docx") == "Untitled.docx"
    assert export._safe_filename("x" * 400, ".pdf").endswith(".pdf")
    assert len(export._safe_filename("x" * 400, ".pdf")) == 255


def test_pdf_text_maps_quotes_and_drops_unsupported_accents_without_black_boxes():
    assert export._pdf_text("“Hi” ’s") == '"Hi" \'s'
    assert export._pdf_text("Olá ọlá") == "Olá olá".replace("olá", "olá")
    assert export._pdf_text("你") == "?"  # no base letter at all
    assert export._pdf_text("a\x00b") == "ab"


def test_pdf_renders_headings_bullets_and_the_legal_banner():
    sections = [{"heading": "Terms & Scope", "body": "One\n\n- x\n- y"}]
    data = export.render_pdf(_doc(sections, title="Deal (draft)"), DISCLAIMER).decode("latin-1")
    assert data.startswith("%PDF")
    assert "(Terms & Scope)" in data and "(Deal \\(draft\\))" in data
    assert data.count("not legal advice") >= 1
    plain = export.render_pdf(_doc(sections), None).decode("latin-1")
    assert "not legal advice" not in plain


def test_pdf_paginates_long_documents():
    body = "\n\n".join(f"Paragraph {i} " + "word " * 60 for i in range(80))
    data = export.render_pdf(_doc([{"heading": "Long", "body": body}]), None).decode("latin-1")
    assert len(re.findall(r"/Type /Page\b", data)) > 1
    assert "Page 2" in data


def test_docx_keeps_unicode_and_strips_xml_invalid_control_characters():
    sections = [{"heading": "Olá \x0b", "body": "ọlá \x00 world\n\n- item"}]
    raw = export.render_docx(_doc(sections, title="T\x07itle"), None)
    parsed = DocxDocument(io.BytesIO(raw))
    text = [p.text for p in parsed.paragraphs]
    assert "Olá " in text and "ọlá  world" in text and "item" in text
    assert parsed.core_properties.title == "Title"
    assert any(p.style.name == "List Bullet" for p in parsed.paragraphs)


def test_disclaimer_comes_from_the_template_or_the_legal_kind():
    assert disclaimer_for("nda") == DISCLAIMER
    assert disclaimer_for(None, DocumentKind.safe) == DISCLAIMER
    assert disclaimer_for("one_pager") is None
    assert disclaimer_for(None, DocumentKind.custom) is None
    assert disclaimer_for("unknown", DocumentKind.nda) == DISCLAIMER


def test_every_placeholder_in_a_template_has_a_field():
    for tmpl in DOCUMENT_TEMPLATES.values():
        keys = {f.key for f in tmpl.fields}
        used = set(re.findall(r"\{\{(\w+)\}\}", " ".join(s.body for s in tmpl.sections)))
        assert used <= keys, (tmpl.key, used - keys)
        assert all(f.key in used for f in tmpl.fields), tmpl.key  # no orphan fields
        assert all(not (s.body and (s.ai or s.autofill)) for s in tmpl.sections), tmpl.key


def test_render_body_is_single_pass_and_brackets_missing_values():
    nda = DOCUMENT_TEMPLATES["nda"]
    out = render_body(
        nda,
        "{{receiving_party}} / {{disclosing_party}} / {{nope}}",
        {"receiving_party": "{{disclosing_party}}"},
    )
    assert out == "{{disclosing_party}} / [Disclosing party] / [nope]"


def test_instantiate_defaults_fill_the_legal_text():
    _kind, _title, sections = instantiate("nda")
    text = " ".join(s["body"] for s in sections)
    assert "2 years" in text and "[Receiving party]" in text


def test_record_lines_skip_noise_and_handle_nameless_records():
    assert _record_line({"name": "Fees", "est_monthly": 0, "pricing_basis": "1%"}) == (
        "- Fees (pricing basis: 1%)"
    )
    assert _record_line({"model_type": "subscription", "tiers": [{"name": "x"}]}) == (
        "- subscription"
    )
    assert _record_line({"goals": ["a", "b"], "map_x": 0.5}) == "- goals: a, b"
    assert _record_line({"name": " "}) == ""


def test_resolve_spec_handles_missing_sources_and_unknown_specs():
    ctx = Context(
        profile={"name": " "}, canvases={}, records={}, user_name=None, today="2026-10-25"
    )
    assert resolve_spec("profile.name", ctx) is None
    assert resolve_spec("user.full_name", ctx) is None
    assert resolve_spec("canvas.lean.problem", ctx) is None
    assert resolve_spec("record.persona", ctx) is None
    assert resolve_spec("bogus.thing", ctx) is None
    assert resolve_spec("today", ctx) == ("2026-10-25", "system")
    ctx.canvases = {"mission_vision": {"mission": "Move money"}}
    assert resolve_spec("canvas.mission_vision.mission", ctx) == ("Move money", "business_builder")


def test_build_form_for_a_template_with_no_fields_has_no_unfilled_required():
    ctx = Context(profile={}, canvases={}, records={}, user_name=None, today="2026-10-25")
    form = build_form(ctx, DOCUMENT_TEMPLATES["meeting_notes"], {})
    assert form["fields"] == [] and form["unfilled_required"] == []
    assert form["ai_available"] is False and form["disclaimer"] is None
