"""Export a rich document as a stored PDF or DOCX file (gap-list row 12).

The export is an ordinary private `DocumentFile`: same storage, same scan gate, same signed links,
so it can be sent for signature like an upload. `source_document_id` / `source_version` record what
it was made from. Re-exporting an unchanged document returns the file already made.
"""

import io
import re
import unicodedata
import uuid
from html import escape

from docx import Document as DocxDocument
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.core.errors import AppError
from app.db.models.document import Document, DocumentFile
from app.db.models.enums import ScanStatus
from app.platform.events import event_bus
from app.platform.storage import get_storage
from app.services.documents.folders import ensure_path
from app.services.documents.scanner import scan_file
from app.services.documents.template_defs import disclaimer_for

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
FORMATS = {"pdf": (PDF, ".pdf"), "docx": (DOCX, ".docx")}
FORMATS_BY_TYPE = dict(FORMATS.values())
MAX_EXPORT_BYTES = 15 * 1024 * 1024  # the upload limit: an export must be a file we'd accept
MAX_SOURCE_CHARS = 200_000
MAX_LINES = 2000  # renderers are super-linear in line count; bounds the time the row lock is held
_FILENAME_UNSAFE = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')
_BULLET = re.compile(r"^\s*[-*•]\s+(.*)$")

# Type 1 Helvetica is WinAnsi (cp1252) only. Common typographic characters are mapped to their
# plain forms and anything else is reduced to its base letter (an accent is dropped), never
# printed as a black box. DOCX keeps the original text.
_PDF_MAP = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "‑": "-", " ": " "})


def _xml_ok(cp: int) -> bool:
    """XML 1.0 Char: tab, LF, CR and everything above space except surrogates and U+FFFE/FFFF
    (also drops DEL and the C1 controls' neighbour 0x7f)."""
    if cp in (0x09, 0x0A, 0x0D):
        return True
    return (
        0x20 <= cp <= 0xD7FF and cp != 0x7F or 0xE000 <= cp <= 0xFFFD or 0x10000 <= cp <= 0x10FFFF
    )


def _clean(text: str) -> str:
    return "".join(ch for ch in text if _xml_ok(ord(ch)))


def _pdf_text(text: str) -> str:
    out = []
    for ch in _clean(text).translate(_PDF_MAP):
        try:
            ch.encode("cp1252")
            out.append(ch)
        except UnicodeEncodeError:
            base = unicodedata.normalize("NFKD", ch).encode("ascii", "ignore").decode()
            out.append(base or "?")
    return "".join(out)


def blocks(body: str) -> list[tuple[str, str]]:
    """The body as `("p", text)` paragraphs and `("li", text)` bullets. A blank line ends a
    paragraph; single newlines inside one are kept (as line breaks)."""
    out: list[tuple[str, str]] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            out.append(("p", "\n".join(para)))
            para.clear()

    for line in _clean(body).splitlines():
        if not line.strip():
            flush()
        elif m := _BULLET.match(line):
            flush()
            out.append(("li", m.group(1).strip()))
        else:
            para.append(line.rstrip())
    flush()
    return out


def _safe_filename(title: str, ext: str) -> str:
    stem = _FILENAME_UNSAFE.sub(" ", _clean(title)).strip(" .") or "Untitled"
    return " ".join(stem.split())[: 255 - len(ext)] + ext


def _footer(canvas: Canvas, *, label: str, disclaimer: str | None) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.grey)
    canvas.drawString(20 * mm, 12 * mm, _pdf_text(label)[:110])
    canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, f"Page {canvas.getPageNumber()}")
    if disclaimer:
        canvas.drawString(20 * mm, 8 * mm, "Draft, not legal advice. Have a lawyer review it.")
    canvas.restoreState()


def render_pdf(doc: Document, disclaimer: str | None) -> bytes:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=20)
    heading = ParagraphStyle(
        "h", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=13, spaceBefore=12
    )
    body = ParagraphStyle("b", parent=styles["BodyText"], fontName="Helvetica", leading=14)
    item = ParagraphStyle("li", parent=body, leftIndent=14, bulletIndent=4)
    note = ParagraphStyle("n", parent=body, fontName="Helvetica-Oblique", fontSize=9, leading=12)
    label = f"{doc.title or 'Untitled'}  |  version {doc.version}  |  Cofoundaz"
    tpl = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=_pdf_text(doc.title or "Untitled"),
        author="Cofoundaz",
    )

    def p(text: str, style: ParagraphStyle, bullet: str | None = None) -> Paragraph:
        return Paragraph(
            escape(_pdf_text(text), quote=False).replace("\n", "<br/>"), style, bulletText=bullet
        )

    flow: list = [p(doc.title or "Untitled", title)]
    if disclaimer:
        box = Table([[p(disclaimer, note)]], colWidths=[A4[0] - 40 * mm])
        box.setStyle(
            TableStyle(
                [
                    ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#B45309")),
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#FEF3C7")),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        flow += [box, Spacer(1, 6)]
    for sec in doc.sections:
        if sec.get("heading"):
            flow.append(p(sec["heading"], heading))
        for kind, text in blocks(sec.get("body") or ""):
            if kind == "li":
                flow.append(p(text, item, "\u2022"))
            else:
                flow.extend(p(line, body) for line in text.split("\n"))
                flow.append(Spacer(1, 4))
    tpl.build(
        flow,
        onFirstPage=lambda c, _d: _footer(c, label=label, disclaimer=disclaimer),
        onLaterPages=lambda c, _d: _footer(c, label=label, disclaimer=disclaimer),
    )
    return buf.getvalue()


def render_docx(doc: Document, disclaimer: str | None) -> bytes:
    out = DocxDocument()
    out.core_properties.title = _clean(doc.title or "Untitled")[:255]
    out.core_properties.author = "Cofoundaz"
    out.add_heading(_clean(doc.title or "Untitled"), level=0)
    if disclaimer:
        run = out.add_paragraph().add_run(disclaimer)
        run.italic = True
        run.font.size = Pt(9)
        footer = out.sections[0].footer.paragraphs[0]
        footer.text = "Draft, not legal advice. Have a lawyer review it."
    for sec in doc.sections:
        if sec.get("heading"):
            out.add_heading(_clean(sec["heading"]), level=1)
        for kind, text in blocks(sec.get("body") or ""):
            out.add_paragraph(text, style="List Bullet" if kind == "li" else None)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def existing_export(db: Session, doc: Document, content_type: str) -> DocumentFile | None:
    """The live export of this document in this format made from the same version, title and
    folder (a rename or move does not bump the version but changes the file's name and place)."""
    ext = FORMATS_BY_TYPE[content_type]
    return (
        db.query(DocumentFile)
        .filter_by(
            startup_id=doc.startup_id,
            source_document_id=doc.id,
            source_version=doc.version,
            content_type=content_type,
            filename=_safe_filename(doc.title, ext),
            folder=doc.folder,
        )
        .filter(DocumentFile.trashed_at.is_(None))
        .order_by(DocumentFile.created_at.desc())
        .first()
    )


def export_document(
    db: Session, doc: Document, fmt: str, *, actor_id: uuid.UUID
) -> tuple[DocumentFile, bool]:
    """`(file, created)`. The caller holds the document row lock (`get_document(lock=True)`), so
    two simultaneous exports of one version make one file, and has checked the role and that the
    document is not trashed."""
    content_type, ext = FORMATS[fmt]
    found = existing_export(db, doc, content_type)
    if found is not None:
        return found, False
    lines = sum(len(_clean(sec.get("body") or "").splitlines()) for sec in doc.sections)
    if len(doc.search_text or "") > MAX_SOURCE_CHARS or lines > MAX_LINES:
        raise AppError("VALIDATION_ERROR", "This document is too large to export.", 422)
    disclaimer = disclaimer_for(doc.template_key, doc.kind)
    content = (render_pdf if fmt == "pdf" else render_docx)(doc, disclaimer)
    if len(content) > MAX_EXPORT_BYTES:
        raise AppError("VALIDATION_ERROR", "This document is too large to export.", 422)
    storage_key = f"documents/{doc.startup_id}/{uuid.uuid4().hex}{ext}"
    get_storage().save_private(storage_key, content, content_type)
    row = DocumentFile(
        startup_id=doc.startup_id,
        uploaded_by_id=actor_id,
        folder=doc.folder,
        filename=_safe_filename(doc.title, ext),
        content_type=content_type,
        size_bytes=len(content),
        storage_key=storage_key,
        url=None,
        is_private=True,
        scan_status=ScanStatus.pending,
        source_document_id=doc.id,
        source_version=doc.version,
    )
    db.add(row)
    db.flush()
    ensure_path(db, doc.startup_id, doc.folder)
    scan_file(db, row, content)
    event_bus.publish(
        db,
        "document.exported",
        {
            "startup_id": str(doc.startup_id),
            "document_id": str(doc.id),
            "file_id": str(row.id),
            "format": fmt,
        },
    )
    return row, True
